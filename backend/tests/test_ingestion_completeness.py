"""A document is only 'completed' when everything in it was actually extracted and stored."""
from datetime import datetime, timedelta, timezone
from unittest.mock import MagicMock

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from openpyxl import Workbook

from app.api.upload import _process_upload_safely
from app.services.ingestion_service import (
    IngestionService,
    PartialIngestionError,
    UserFacingIngestionError,
    _read_text_file,
)


def _service(db=None):
    embedder = MagicMock()
    embedder.embed_text.side_effect = lambda texts: [[0.0, 1.0] for _ in texts]
    return IngestionService(db=db or MagicMock(), embedder=embedder, lexical_store=MagicMock())


def _counting_db():
    db = MagicMock()
    stored = set()
    db.save_document_chunks.side_effect = lambda records: stored.update(r["chunk_index"] for r in records) or len(records)
    db.count_document_chunks.side_effect = lambda document_id: len(stored)
    return db


def _run(service, path, name="doc.txt"):
    service.process_file_background(str(path), name, "hash", "box-1", "doc-1")


# ── Stored-chunk verification ──

def test_missing_stored_chunks_fail_the_document(tmp_path):
    path = tmp_path / "doc.txt"
    path.write_text("Some policy text that will be chunked.")
    db = MagicMock()
    db.count_document_chunks.return_value = 0  # saves "succeeded" but nothing landed

    with pytest.raises(PartialIngestionError, match="expected 1, found 0"):
        _run(_service(db), path)

    db.delete_chunks_by_document.assert_called_once_with("doc-1")
    db.update_document_status.assert_not_called()


def test_index_invalidation_failure_keeps_completed_document(tmp_path):
    path = tmp_path / "doc.txt"
    path.write_text("Some policy text that will be chunked.")
    service = _service(_counting_db())
    service.lexical_store.invalidate_box.side_effect = RuntimeError("index busy")

    _run(service, path)

    service.db.update_document_status.assert_called_once_with("doc-1", "completed")
    service.db.delete_chunks_by_document.assert_not_called()


def test_query_cache_failure_does_not_mark_document_failed(monkeypatch):
    from app.core import cache
    service = MagicMock()
    monkeypatch.setattr(cache.global_query_cache, "invalidate_box", MagicMock(side_effect=RuntimeError("boom")))

    _process_upload_safely(service, "/tmp/x.txt", "x.txt", "hash", "box-1", "doc-1")

    service.db.update_document_status.assert_not_called()


# ── Partly scanned PDFs ──

class _Page:
    def __init__(self, text):
        self.text = text

    def get_text(self, mode):
        return [(10, 10, 500, 40, self.text, 0, 0)] if self.text else []


class _Pdf:
    def __init__(self, texts):
        self.pages = [_Page(t) for t in texts]
        self.page_count = len(texts)

    def __enter__(self):
        return self

    def __exit__(self, *exc):
        return False

    def get_toc(self):
        return []

    def load_page(self, i):
        return self.pages[i]


def _pdf_pages(monkeypatch, texts):
    import fitz
    monkeypatch.setattr(fitz, "open", lambda *a, **k: _Pdf(texts))


def test_pdf_with_blank_pages_is_accepted(monkeypatch):
    # Blank or image-only pages are normal in user uploads; only an all-empty PDF is rejected.
    _pdf_pages(monkeypatch, ["", "", "Real content on this page. " * 20, ""])

    saved = _service(_counting_db())._process_pdf_streaming("x.pdf", "x.pdf", "hash", "box-1", "doc-1")

    assert saved > 0


def test_words_wrapped_in_bidi_marks_keep_their_spaces(monkeypatch):
    # Google Docs PDF exports wrap each word in LRO…PDF (U+202D…U+202C) instead of spaces
    words = "the fundamental rights are termed fundamental for a reason".split()
    _pdf_pages(monkeypatch, ["".join(f"‭{w}‬" for w in words) * 10])
    saved = []
    db = MagicMock()
    db.save_document_chunks.side_effect = lambda records: saved.extend(records) or len(records)

    _service(db)._process_pdf_streaming("x.pdf", "x.pdf", "hash", "box-1", "doc-1")

    content = " ".join(r["content"] for r in saved)
    assert "rights are termed fundamental for a reason" in content
    assert "‭" not in content and "‬" not in content


def test_invisible_characters_are_cleaned():
    from app.services.ingestion_service import _clean_extracted_text
    assert _clean_extracted_text("‭Part‬‭III‬") == " Part III "
    assert _clean_extracted_text("consti­tution​") == "constitution"


# ── Abandoned 'processing' documents ──

def test_stale_processing_detection():
    service = _service()
    old = (datetime.now(timezone.utc) - timedelta(minutes=120)).isoformat()
    recent = datetime.now(timezone.utc).isoformat()

    assert service.is_stale_processing({"status": "processing", "created_at": old})
    assert not service.is_stale_processing({"status": "processing", "created_at": recent})
    assert not service.is_stale_processing({"status": "completed", "created_at": old})
    assert not service.is_stale_processing({"status": "processing", "created_at": None})


def test_fail_stale_documents_uses_configured_cutoff():
    db = MagicMock()
    db.fail_stale_processing_documents.return_value = 2
    service = IngestionService(db=db, embedder=MagicMock(), lexical_store=MagicMock(), stale_processing_minutes=15)

    assert service.fail_stale_documents("box-1") == 2
    minutes, message, box_id = db.fail_stale_processing_documents.call_args.args
    assert (minutes, box_id) == (15, "box-1")
    assert "upload the document again" in message


# ── Text encodings ──

def test_utf16_text_is_decoded(tmp_path):
    path = tmp_path / "u16.txt"
    path.write_bytes("Gehalt über 5 € pro Stunde".encode("utf-16"))
    assert _read_text_file(str(path)) == "Gehalt über 5 € pro Stunde"


def test_windows_1252_text_is_decoded(tmp_path):
    path = tmp_path / "cp.csv"
    path.write_bytes("name,amount\nJosé,€5".encode("cp1252"))
    assert _read_text_file(str(path)) == "name,amount\nJosé,€5"


def test_utf8_bom_is_stripped(tmp_path):
    path = tmp_path / "bom.csv"
    path.write_bytes("﻿name,amount\nA,1".encode("utf-8"))
    assert _service()._extract_csv(str(path)) == "name: A | amount: 1"


# ── DOCX coverage ──

def test_docx_reads_content_controls_tracked_insertions_and_headers(tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Plain paragraph.")
    body = doc.element.body
    body.insert(len(body) - 1, parse_xml(
        f'<w:sdt {nsdecls("w")}><w:sdtContent><w:p><w:r><w:t>Inside a content control.</w:t></w:r></w:p>'
        f'</w:sdtContent></w:sdt>'
    ))
    body.insert(len(body) - 1, parse_xml(
        f'<w:p {nsdecls("w")}><w:ins w:id="1" w:author="a"><w:r><w:t>Tracked insertion.</w:t></w:r></w:ins>'
        f'<w:del w:id="2" w:author="a"><w:r><w:delText>Deleted text.</w:delText></w:r></w:del></w:p>'
    ))
    doc.sections[0].header.paragraphs[0].text = "Policy No. HR-7"
    path = tmp_path / "rich.docx"
    doc.save(path)

    text = _service()._extract_docx(str(path))

    assert "Plain paragraph." in text
    assert "Inside a content control." in text
    assert "Tracked insertion." in text
    assert "Deleted text." not in text
    assert "## Headers and Footers" in text and "Policy No. HR-7" in text


# ── XLSX ──

def test_xlsx_formula_without_saved_value_is_rejected(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Item", "Price", "Total"])
    ws.append(["Pen", 2, "=B2*10"])  # openpyxl writes no cached value
    path = tmp_path / "formulas.xlsx"
    wb.save(path)

    with pytest.raises(UserFacingIngestionError, match="1 formula cells"):
        _service()._extract_xlsx(str(path))


def test_xlsx_row_wider_than_header_keeps_extra_cells(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Item"])
    ws.append(["Pen", "blue"])
    path = tmp_path / "wide.xlsx"
    wb.save(path)

    text = _service()._extract_xlsx(str(path))

    assert "Item: Pen | Col1: blue" in text
